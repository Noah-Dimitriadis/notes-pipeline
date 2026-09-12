from __future__ import annotations

import re
import shutil
import subprocess
import warnings
from collections import Counter
from pathlib import Path

from ..models import Deck, Slide


def extract(path: Path) -> Deck:
    """Extract a Deck from a .pptx or .pdf slide export."""
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    raise ValueError(f"Unsupported deck format: {suffix!r} (expected .pptx or .pdf)")


# -- .pptx ------------------------------------------------------------------


def _extract_pptx(path: Path) -> Deck:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[Slide] = []

    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_shape_id = title_shape.shape_id if title_shape is not None else None
        title: str | None = None
        if title_shape is not None and title_shape.has_text_frame:
            text = title_shape.text_frame.text.strip()
            title = text or None

        # Gotcha: python-pptx yields shapes in XML order, not visual order —
        # sort by (top, left) or multi-column body text comes out scrambled.
        # Compare by shape_id, not identity: `slide.shapes.title` and the
        # `slide.shapes` iterator hand back separate wrapper objects for the
        # same underlying shape, so `is not title_shape` never matches.
        text_shapes = [
            shape
            for shape in slide.shapes
            if shape.shape_id != title_shape_id
            and getattr(shape, "has_text_frame", False)
            and shape.text_frame.text.strip()
        ]
        text_shapes.sort(key=lambda s: (s.top if s.top is not None else 0, s.left if s.left is not None else 0))
        body_texts = [shape.text_frame.text.strip() for shape in text_shapes]

        if title is None and body_texts:
            # Empty title placeholders are common — fall back to the first
            # line of body text, and don't duplicate it into the body.
            first_line, _, remainder = body_texts[0].partition("\n")
            title = first_line.strip()
            remainder = remainder.strip()
            body_texts[0] = remainder
            body_texts = [t for t in body_texts if t]

        body = "\n\n".join(body_texts)

        notes: str | None = None
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            notes = notes_text or None

        slides.append(Slide(index=i, title=title, body=body, notes=notes))

    return Deck(slides=slides, source=path)


# -- .pdf ---------------------------------------------------------------------


def pdf_to_text(path: Path) -> str:
    """Extract raw text from a PDF via `pdftotext -layout` (poppler), falling
    back to pypdf with a warning if poppler isn't on PATH."""
    if shutil.which("pdftotext") is not None:
        return _pdftotext_layout(path)
    warnings.warn(
        "poppler's `pdftotext` was not found on PATH; falling back to pypdf, "
        "which does not preserve layout as reliably. "
        "Install poppler (`brew install poppler`) for better extraction.",
        RuntimeWarning,
        stacklevel=2,
    )
    return _pypdf_text(path)


def _extract_pdf(path: Path) -> Deck:
    raw_text = pdf_to_text(path)
    pages = raw_text.split("\f")
    if len(pages) > 1 and pages[-1] == "":
        pages = pages[:-1]

    header = _detect_header([page.splitlines() for page in pages])

    slides: list[Slide] = []
    for i, page in enumerate(pages, start=1):
        title, body = _parse_pdf_page(page, header)
        slides.append(Slide(index=i, title=title, body=body, notes=None))

    return Deck(slides=slides, source=path)


def _pdftotext_layout(path: Path) -> str:
    result = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        capture_output=True,
        text=True,
        check=True,
    )
    return result.stdout


def _pypdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\f".join(page.extract_text() or "" for page in reader.pages)


def _normalize_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip())


def _detect_header(pages_lines: list[list[str]]) -> str | None:
    """The running header (course/title/institution line) repeats as the
    first non-empty line on most pages — detect it so it can be stripped.
    `-layout` pads the gap between the course name and institution name
    with a slightly different number of spaces per page (rendering-width
    rounding), so compare with internal whitespace collapsed."""
    firsts = []
    for lines in pages_lines:
        for line in lines:
            if line.strip():
                firsts.append(_normalize_whitespace(line))
                break
    if not firsts:
        return None
    candidate, freq = Counter(firsts).most_common(1)[0]
    if freq >= max(2, len(firsts) // 2):
        return candidate
    return None


def _parse_pdf_page(raw_page: str, header: str | None) -> tuple[str | None, str]:
    """Title = first non-empty line after stripping the running header; the
    trailing slide-number line is stripped from the end. Column layout in
    between (preserved by `-layout`) passes through untouched."""
    lines = raw_page.splitlines()
    nonempty = [i for i, line in enumerate(lines) if line.strip()]
    if not nonempty:
        return None, ""

    start, end = 0, len(nonempty)
    if header is not None and _normalize_whitespace(lines[nonempty[start]]) == header:
        start += 1
    if end > start and lines[nonempty[end - 1]].strip().isdigit():
        end -= 1

    content = nonempty[start:end]
    if not content:
        return None, ""

    title_line = content[0]
    body_end_line = nonempty[end] if end < len(nonempty) else len(lines)
    body = "\n".join(lines[title_line + 1 : body_end_line]).strip("\n")
    return lines[title_line].strip(), body
